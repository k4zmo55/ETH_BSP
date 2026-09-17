# -*- coding: utf-8 -*-
"""E-AETIS staj sunumu - slayt slayt konusma metni (Word / .docx) uretici."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))

NAVY = RGBColor(0x00, 0x46, 0x94)
INK = RGBColor(0x10, 0x18, 0x26)
MUTED = RGBColor(0x5b, 0x6b, 0x80)

SLIDES = [
    (1, "Kapak",
     "Merhaba, ben Mehmet Akif Seçkin. Staj boyunca EHSİM'de geliştirdiğim E-AETIS projesini "
     "anlatacağım: sıfırdan yazılmış, taşınabilir bir Ethernet sürücüsü ve bunu test eden bir "
     "masaüstü arayüz. On dakikada projenin ne olduğunu, nasıl çalıştığını ve neyi doğruladığımı "
     "anlatacağım, sonunda sorularınızı alacağım."),

    (2, "Problem",
     "Her ağa bağlı cihazın bir ortak dile ihtiyacı var: mikrodenetleyici üzerinde çalışan uygulama, "
     "ürettiği veriyi ağın anlayacağı Ethernet paketine çeviren bir yazılım katmanına muhtaç — buna "
     "Ethernet sürücüsü diyoruz. Piyasada bunun hazır çözümleri var: STM32'nin kendi HAL_ETH'i, LwIP "
     "TCP/IP yığını ve genelde bunun üstünde bir RTOS. Ama bu üçlü ağırdır ve her karta doğrudan "
     "taşınmaz; her STM32 ailesinin HAL'i farklı register haritası ve farklı API kullanır. Benim "
     "görevim buydu: hafif, taşınabilir bir sürücüyü sıfırdan, register seviyesinde yazmak."),

    (3, "Ne Yaptım",
     "İki parçadan oluşan bir sistem geliştirdim. Kartın içinde, C ile yazılmış Ethernet_BSP "
     "kütüphanesi var — register seviyesinde MAC, DMA ve PHY erişimi yapıyor, RTOS yok, LwIP yok, "
     "ARP/ICMP/UDP'yi elle yazdım, isteğe bağlı bir IAP bootloader ekledim. Bilgisayar tarafında ise "
     "Python/PyQt5 ile yazdığım E-AETIS GUI'si var; kartla UDP üzerinden konuşuyor, telemetri okuyor, "
     "performans ve jitter ölçüyor, hata enjeksiyonuyla fuzz testi yapıyor ve ağdan firmware "
     "güncelliyor. İkisi aynı protokolle konuşuyor."),

    (4, "Tasarım Felsefesi — HAL_ETH / LwIP / FreeRTOS neden kullanılmadı",
     "HAL_ETH, LwIP ve FreeRTOS'u bilinçli olarak kullanmadım çünkü amacım ürünü en hızlı şekilde "
     "piyasaya sürmek değil, MAC, PHY ve DMA seviyesinde gerçekte ne olup bittiğini görebileceğim, "
     "öngörülebilir bir sistem kurmaktı; hazır yığın kullansaydım her register erişimi üç ayrı "
     "kütüphanenin arkasında gizlenecek, LwIP'in dinamik pbuf havuzu yüzünden bellek kullanımı "
     "öngörülemez hale gelecek, FreeRTOS'un görevler arası geçişleri zamanlamaya jitter katacak ve "
     "HAL_ETH'in aileler arası farklı API'leri yüzünden bir karttan diğerine taşınmak da kolay "
     "olmayacaktı — bunun karşılığında elde ettiğim şey statik bellek, tek bağlamda deterministik "
     "zamanlama ve tek bir config dosyasıyla dört farklı MCU'ya taşınabilen bir mimari oldu; tabii "
     "bunun bedeli de var, TCP/DHCP/DNS gibi hazır yığının sağladığı özellikleri kendim yazmadım, "
     "sadece ihtiyacım olan ARP/ICMP/UDP'yi elle yazdım."),

    (5, "Mimari",
     "Mimarinin kalbinde üç bağımsız eksen var: hangi MCU/MAC ailesinin kullanılacağı, hangi PHY "
     "çipinin kullanılacağı ve donanımdan tamamen bağımsız çekirdek mantığı. Kullanıcı sadece "
     "eth_config.h'yi düzenliyor; eth_device.h bunu doğru MAC ailesine eşliyor. MAC ailesi için "
     "21 fonksiyonluk bir port sözleşmesi, PHY için 4 fonksiyonluk bir sözleşme tanımladım — her "
     "yeni donanım sadece bu sözleşmeyi implemente ediyor, çekirdek koda (eth_driver.c, eth_app.c) "
     "hiç dokunmuyor. Bu ayrım sayesinde bir eksende yapılan değişiklik diğerini etkilemiyor."),

    (6, "Taşınabilirlik",
     "Bunun sonucu somut: aynı kütüphane, tek bir config dosyasıyla dört farklı STM32'de çalışıyor — "
     "Cortex-M33, iki farklı Cortex-M7 ve bir Cortex-M4, iki farklı MAC ailesi (Synopsys EQOS ve "
     "klasik GMAC), üç farklı PHY çipiyle. Yeni bir MCU eklemek eth_device.h'ye birkaç satır eklemek "
     "demek; port ve PHY dosyalarına hiç dokunulmuyor. MAC ile PHY arasındaki RMII arayüzünün 50 "
     "MHz'lik ortak bir referans saatine ihtiyacı var — bu saat olmadan ikisi birbirini "
     "örnekleyemiyor, bu yüzden kartı açtığımda ilk kontrol ettiğim şey hep bu saat."),

    (7, "Çekirdek Sürücü",
     "Çekirdekte, DMA ile CPU'nun paylaştığı 8 elemanlı bir alım halkası var. Üç kural bunu güvenli "
     "tutuyor: zero-copy, yani DMA'nın yazdığı tampon üzerinden doğrudan işliyoruz, ara kopyalama "
     "yok; her descriptor'da bir OWN biti var, o an DMA'nın mı CPU'nun mu sahipliğinde olduğunu "
     "söylüyor, ikisi asla aynı anda aynı elemana dokunmuyor; ve bu bölgeyi MPU ile non-cacheable "
     "işaretledim ki CPU, DMA'nın yazdığı veriyi cache'ten değil doğrudan bellekten okusun."),

    (8, "Paket İşleme Akışı",
     "Gelen her çerçeve sırayla süzülüyor: önce EtherType'a bakılıp ARP mı IPv4 mü olduğuna karar "
     "veriliyor, IPv4 ise protokol alanına bakılıp ICMP mi UDP mi ayrılıyor, UDP ise hedef port'a "
     "göre 5000 numaralı port E-AETIS komutlarına, diğerleri kullanıcının kendi callback'ine "
     "yönlendiriliyor. Bilinçli bir kısıt var: IP parçalanmış paketleri (fragment) yeniden "
     "birleştirmiyorum, bunlar sessizce düşürülüyor — bu bir hata değil, tasarım kararı; hata "
     "enjeksiyonu testinde gönderdiğim 2048 baytlık paket de tam olarak bu yüzden zaman aşımına "
     "uğruyor."),

    (9, "IAP Bootloader",
     "Ağdan firmware güncellemek için START_IAP, FW_DATA, END_IAP adımlarından oluşan basit bir "
     "protokol yazdım; her parça CRC32 ile doğrulanıyor. Burada en çok zorlandığım yer, yeni "
     "firmware'e atlama anıydı: ilk denemede kart anında HardFault'a düşüyordu, çünkü atlarken hâlâ "
     "çalışan Ethernet DMA'sı ve SysTick kesmeleri, yeni uygulamanın henüz kurmadığı bir vektör "
     "tablosuna düşüyordu. Çözüm, atlamadan hemen önce Ethernet'i durdurup tüm kesmeleri kapatmak "
     "oldu. Önemli bir güvenlik notu: bu mekanizmada kimlik doğrulama yok, sadece bütünlük kontrolü "
     "var, bu yüzden varsayılan olarak kapalı tutuyorum. Sağdaki ekran görüntüsü bunun gerçek bir "
     "çalıştırmasını gösteriyor: küçük bir LED test firmware'i yüklendi, kart tüm parçaları alıp "
     "CRC32'yi doğruladı ve FLASH_SUCCESS|JUMP_OK ile yeni uygulamaya atladı — LED beklendiği gibi "
     "yanıp söndü."),

    (10, "Test Arayüzü (GUI)",
     "Geliştirdiğim arayüzün yedi sekmesi var: kart değişkenlerini otomatik keşfedip izleyen bir "
     "sekme, ping ile gecikme ölçen bir sekme, serbest komut için bir UDP konsolu, PHY/DMA "
     "register'larını gösteren bir teşhis ekranı, performans ve jitter ölçen bir sekme, hata "
     "enjeksiyonu ve bootloader sekmesi. Kart yeni bir değişken kaydettiğinde arayüz onu elle "
     "tanımlamaya gerek kalmadan otomatik buluyor. Soldaki ekran görüntüsü gerçek bir kayıttan: "
     "kart açılınca beş değişkeni (sıcaklık, besleme gerilimi, çalışma süresi, nem, led) kendisi "
     "keşfetti, değerleri canlı grafikte izliyorum ve olay günlüğünde kartın 0.39 milisaniyede "
     "yanıt verdiğini görüyorum — yani bu, gerçekten karttan canlı veri okuyan çalışan bir sistem."),

    (11, "Süreç",
     "Arayüz de sürücü gibi gün gün büyüdü; ilk günlerdeki boş pencereden, bugünkü canlı performans "
     "ve hata enjeksiyonu ekranlarına kadar her yeni test yeteneği kartta çalışır hale geldikçe "
     "eklendi."),

    (12, "Doğrulama & Sonuçlar",
     "Çalışıyor gibi görünmesi bana yetmedi, her iddiayı donanımda ölçtüm: MDIO register alanlarını, "
     "RMII seçim kodunu, DMA tail pointer'ını, MPU alan yerleşimini referans kılavuzundan tek tek "
     "doğruladım. Sonuç olarak ortalama 0.44 milisaniye gecikme, yüzde sıfır paket kaybı ölçtüm ve "
     "hata enjeksiyonu sırasında kart hiç çökmedi."),

    (13, "Entegrasyon — Başka Bir Kullanıcı Bu Projeyi Nasıl Kullanır?",
     "Bu kütüphaneyi başka bir geliştirici de kolayca kullanabilir: Inc ve Src klasörlerini "
     "projesine kopyalar, eth_config.h'de MCU/PHY/IP'sini seçer (soldaki kod), main.c'de "
     "ETH_BSP_Init ve ProcessEvents'i çağırır (sağ üstteki kod). Yeni bir MCU eklemek istersen "
     "eth_device.h'ye tek bir #elif bloğu eklemen yeterli (sol alt); yeni bir PHY çipi eklemek "
     "istersen de eth_phy.h'deki dört fonksiyonu (Bringup, GetSpeedDuplex, SetLoopback, GetName) "
     "yeni bir dosyada implemente etmen yeterli (sağ alt) — ikisinde de port/PHY sözleşmesi "
     "sayesinde çekirdek koda hiç dokunulmuyor. Detaylı adımlar ve daha fazla örnek README'de."),

    (14, "Araçlar & Teknolojiler",
     "Bu projede sadece C kodu yazmadım; CMake, Git, GCC/Ninja, STM32CubeMX gibi araçlarla bir "
     "geliştirme iş akışı kurdum, DMA, Ethernet, I2C, SPI, UART, kesme yönetimi gibi donanım "
     "protokollerini de pratikte öğrendim — bazılarını, mesela FreeRTOS temellerini ve I2C/SPI'yi, "
     "bu projede kullanmasam da staj boyunca ayrıca çalıştım."),

    (15, "Kazanımlar",
     "Kırk günlük staj boyunca gereksinimden mimariye, uygulamadan donanımda doğrulamaya uzanan tam "
     "bir mühendislik döngüsünü kendi ellerimle tamamladım — dört farklı MCU'yu destekleyen, sıfır "
     "malloc çağrısı yapan, 21 fonksiyonluk bir port sözleşmesiyle taşınabilir bir kütüphane ortaya "
     "çıktı. Bana kattığı en önemli şey, savunma sanayiini yakından tanımak ve bir sistemi sıfırdan "
     "yazma özgüveniydi. Sorularınızı almaktan memnuniyet duyarım."),
]


def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


doc = Document()

section = doc.sections[0]
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.2)
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.8)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("E-AETIS Staj Sunumu — Konuşma Metni")
run.bold = True
run.font.size = Pt(20)
run.font.color.rgb = NAVY

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = sub.add_run("15 slayt için sayfa sayfa konuşma notları · EAETIS_Staj_Sunumu_Taslak.pptx ile birlikte kullanılır")
run.italic = True
run.font.size = Pt(10.5)
run.font.color.rgb = MUTED

doc.add_paragraph()

for idx, name, text in SLIDES:
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(2.0)
    table.columns[1].width = Cm(13.6)

    cell0 = table.cell(0, 0)
    set_cell_shading(cell0, "10233F")
    p0 = cell0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = p0.add_run("%02d" % idx)
    r0.bold = True
    r0.font.size = Pt(20)
    r0.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    cell1 = table.cell(0, 1)
    set_cell_shading(cell1, "EEF1F6")
    p1 = cell1.paragraphs[0]
    r1 = p1.add_run(name)
    r1.bold = True
    r1.font.size = Pt(13)
    r1.font.color.rgb = INK

    body = doc.add_paragraph()
    body.paragraph_format.space_before = Pt(8)
    body.paragraph_format.space_after = Pt(18)
    rb = body.add_run(text)
    rb.font.size = Pt(11.5)
    rb.font.color.rgb = INK

out_path = os.path.join(HERE, "EAETIS_Sunum_Konusma_Metni.docx")
doc.save(out_path)
print("OK:", out_path)
